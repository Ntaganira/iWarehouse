"""Convert the High Bytes product/price workbook into the iWarehouse import template.

The source sheet is a human-maintained price list: repeated glass products with
different sheet sizes, accessory rows whose "Size" is really a pack/volume, and
measurement-only continuation rows such as "1*25YARDS" or "6.38mm".

The target is the 14-column iWarehouse template (Type, Code, Glass type, ...).
The requested generic target fields map as follows:

    Product Title / Name  -> Glass type + Colour/finish + Thickness, or Accessory name
    Dimensions / Size     -> Sheet width/height (glass), Notes (accessories)
    Base Cost Price       -> lower bound is kept in Notes
    Max/Selling Price     -> Sheet price / Unit price (the upper bound)

Run:
    python tools/convert_high_bytes_products.py
"""

from __future__ import annotations

import argparse
import math
import re
import shutil
import warnings
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd

warnings.filterwarnings(
    "ignore",
    message="Workbook contains no default style, apply openpyxl's default",
    category=UserWarning,
)

SOURCE_WORKBOOK = Path(r"C:\Users\Muganga\Downloads\High bytes products and prices.xlsx")
TARGET_TEMPLATE = Path(r"C:\Users\Muganga\Downloads\iwarehouse-product-import.xlsx")
OUTPUT_WORKBOOK = TARGET_TEMPLATE.with_name(TARGET_TEMPLATE.stem + "-filled.xlsx")

GLASS_TYPES = {"CLEAR", "TINTED", "REFLECTIVE", "FROSTED", "TEMPERED", "LAMINATED", "MIRROR"}
ACCESSORY_UNITS = {"M2", "METRE", "PIECE", "ROLL", "TUBE", "KG", "LITRE", "SET", "BOX", "PACK", "BAG"}
TARGET_COLUMNS = (
    "Type", "Code", "Glass type", "Colour / finish / pattern", "Thickness (mm)",
    "Accessory name", "Accessory unit", "Sheet width (mm)", "Sheet height (mm)",
    "Sheet price (RWF)", "Unit price (RWF)", "Reorder level (m2)", "Max level (m2)", "Notes",
)

# Source workbook headers (the names vary between exports).
NAME_ALIASES = {"PRODUCTNAME", "ITEMNAME", "PRODUCTTITLENAME", "NAME"}
SIZE_ALIASES = {"SIZE", "DIMENSIONSSIZE", "DIMENSIONS"}
PRICE_ALIASES = {"PRICERANGE", "PRICESRWF", "PRICE", "SELLINGPRICE"}

CONTINUATION_PATTERNS = (
    re.compile(r"^\d+(?:\.\d+)?(?:\s*[*xX]\s*\d+(?:\.\d+)?)+(?:\s*(?:YARDS?|YDS?))?$", re.I),
    re.compile(r"^\d+(?:\.\d+)?\s*MM(?:\s*[*xX]\s*\d+(?:\.\d+)?)*$", re.I),
    re.compile(r"^\d+(?:\.\d+)?\s*(?:ML|L)(?:\s*\(.*\))?$", re.I),
    re.compile(r"^\d+\s*&\s*[A-Z0-9]+\s*\)?$", re.I),
)

ACCESSORY_WORDS = (
    "TAPE", "SEALANT", "SILICONE", "SILICON", "GUN", "FILE", "RUBBER", "BRACKET",
    "DISC", "DISQUE", "P38", "INDASSA", "SAND", "GLUE", "ADHESIVE", "CEMENT",
    "PUTTY", "FOAM", "WEATHER", "EASY SEAL", "TRITON", "BOOP", "THINNER", "SOLVENT",
)
GLASS_WORDS = (
    "GLASS", "FLOAT", "REFLECT", "REFL", "LAMINATED", "LAMINAT", "MIRROR",
    "TEMPER", "FROST", "ETCH", "TINT", "CLEAR", "BRONZE", "AMBER", "BLUE",
    "GREY", "GRAY", "GREEN", "LAKE", "MILK",
)


@dataclass
class SourceItem:
    excel_row: int
    name: str
    size: str
    price_text: str
    continuations: list[str] = field(default_factory=list)
    kind: str = ""
    glass_type: str = ""
    variant: str = ""
    thickness: float | None = None
    width: int | None = None
    height: int | None = None
    price_min: int | None = None
    price_max: int | None = None
    chosen_price: int | None = None
    suspicious_range: bool = False
    target_name: str = ""
    unit: str = ""
    notes: str = ""


def clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    text = str(value).replace("\u00a0", " ").strip()
    return re.sub(r"\s+", " ", text)


def normalized(value: object) -> str:
    return re.sub(r"[^A-Z0-9]", "", clean_text(value).upper())


def find_column(columns: list[str], aliases: set[str], label: str) -> str:
    for column in columns:
        if normalized(column) in aliases:
            return column
    raise SystemExit(f"Source workbook has no {label} column. Found: {columns}")


def is_continuation(text: str) -> bool:
    return any(pattern.fullmatch(text.strip()) for pattern in CONTINUATION_PATTERNS)


def looks_like_accessory(name: str) -> bool:
    upper = name.upper()
    return any(word in upper for word in ACCESSORY_WORDS)


def looks_like_glass(name: str, size: str) -> bool:
    upper = name.upper()
    if looks_like_accessory(name):
        return False
    if re.search(r"\d+(?:\.\d+)?\s*MM\b", upper):
        return True
    has_glass_word = any(word in upper for word in GLASS_WORDS)
    has_dimensions = bool(re.fullmatch(r"\s*\d+(?:\.\d+)?\s*[*xX]\s*\d+(?:\.\d+)?\s*", size or ""))
    return has_glass_word and has_dimensions


def glass_type(name: str) -> str:
    upper = name.upper()
    if "LAMINAT" in upper:
        return "LAMINATED"
    if "REFLECT" in upper or re.search(r"\bREFL\b", upper):
        return "REFLECTIVE"
    if "MIRROR" in upper:
        return "MIRROR"
    if "TEMPER" in upper:
        return "TEMPERED"
    if "FROST" in upper or "ETCH" in upper:
        return "FROSTED"
    if "TINT" in upper:
        return "TINTED"
    if "CLEAR" in upper:
        return "CLEAR"
    return "TINTED"


def thickness_mm(name: str) -> float | None:
    match = re.search(r"(\d+(?:\.\d+)?)\s*MM\b", name.upper())
    return float(match.group(1)) if match else None


def clean_variant(name: str) -> str:
    variant = name.upper()
    variant = re.sub(r"\d+(?:\.\d+)?\s*MM\b", " ", variant)
    for word in ("GLASS", "FLOAT", "LAMINATED", "REFLECTIVE", "REFLECT", "REFL", "MIRROR",
                 "TEMPERED", "FROSTED", "TINTED", "CLEAR"):
        variant = re.sub(rf"\b{re.escape(word)}\b", " ", variant)
    variant = re.sub(r"\b\d+(?:\.\d+)?\b", " ", variant)
    variant = re.sub(r"[^A-Z0-9 ]+", " ", variant)
    variant = re.sub(r"\s+", " ", variant).strip()
    return variant.title()


def parse_dimensions(size: str) -> tuple[int, int] | None:
    match = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*[*xX]\s*(\d+(?:\.\d+)?)\s*", size or "")
    if not match:
        return None
    width, height = float(match.group(1)), float(match.group(2))
    if width <= 400 and height <= 400:
        width, height = width * 10, height * 10
    elif width <= 400:
        width *= 10
    elif height <= 400:
        height *= 10
    width_i, height_i = int(round(width)), int(round(height))
    if not (1 <= width_i <= 10_000 and 1 <= height_i <= 10_000):
        return None
    return width_i, height_i


def parse_price(text: str) -> tuple[int, int, bool] | None:
    numbers = re.findall(r"\d[\d,]*(?:\.\d+)?", text or "")
    if not numbers:
        return None
    values = [int(round(float(number.replace(",", "")))) for number in numbers]
    low, high = min(values), max(values)
    if low <= 0 or high <= 0:
        return None
    # A range such as "70,000 - 750,000" is almost certainly a typo; keep the
    # lower bound and report it instead of corrupting the selling price.
    suspicious = low > 0 and high > low * 5
    return low, high, suspicious


def accessory_name(base: str, item: SourceItem, index: int, count: int) -> str:
    base = re.sub(r"[\s\-_(&]+$", "", base).strip() or "Accessory"
    if count <= 1:
        return base[:60]
    suffix = item.size or f"#{index + 1}"
    room = max(1, 60 - len(suffix) - 3)
    return f"{base[:room].rstrip()} ({suffix})"[:60]


def accessory_unit(name: str) -> str:
    upper = name.upper()
    if "TAPE" in upper:
        return "ROLL"
    if any(word in upper for word in ("SEALANT", "SILICONE", "SILICON", "GLUE", "ADHESIVE",
                                      "CEMENT", "TRITON", "EASY SEAL")):
        return "TUBE"
    if "RUBBER" in upper:
        return "METRE"
    return "PIECE"


def note_text(item: SourceItem, group_size: int = 1) -> str:
    parts: list[str] = []
    if item.kind == "GLASS":
        parts.append(f"Source: {item.name}")
    if item.size:
        parts.append(f"Size: {item.size}")
    if item.price_text:
        parts.append(f"Price: {item.price_text}")
    if item.continuations:
        parts.append("Continuation: " + " | ".join(item.continuations))
    if item.suspicious_range:
        parts.append("Suspicious range; lower bound used")
    if group_size > 1:
        parts.append(f"Largest of {group_size} sheets")
    return " | ".join(parts)[:255]


def read_source(path: Path) -> tuple[list[SourceItem], int, list[tuple[int, str, str]]]:
    frame = pd.read_excel(path, sheet_name=0, dtype=object, keep_default_na=False)
    frame.columns = [clean_text(column) for column in frame.columns]
    columns = list(frame.columns)
    name_column = find_column(columns, NAME_ALIASES, "Product Name / Item Name")
    size_column = find_column(columns, SIZE_ALIASES, "Size")
    price_column = find_column(columns, PRICE_ALIASES, "Price Range")

    items: list[SourceItem] = []
    skipped: list[tuple[int, str, str]] = []
    attached_continuations = 0

    for index, row in frame.iterrows():
        excel_row = int(index) + 2
        name = clean_text(row[name_column])
        size = clean_text(row[size_column])
        price = clean_text(row[price_column])
        if not name and not size and not price:
            continue

        if name and not size and not price and is_continuation(name):
            if items:
                items[-1].continuations.append(name)
                attached_continuations += 1
            else:
                skipped.append((excel_row, name, "continuation row with no preceding item"))
            continue

        items.append(SourceItem(excel_row=excel_row, name=name, size=size, price_text=price))

    # Classify and parse only after every continuation has been attached, so a
    # missing thickness can be recovered from the following measurement row.
    for item in items:
        if looks_like_glass(item.name, item.size):
            item.kind = "GLASS"
            item.thickness = thickness_mm(item.name)
            if item.thickness is None:
                item.thickness = next(
                    (value for continuation in item.continuations
                     if (value := thickness_mm(continuation)) is not None),
                    None,
                )
            item.width, item.height = parse_dimensions(item.size) or (None, None)
            item.glass_type = glass_type(item.name)
            item.variant = clean_variant(item.name)
        else:
            item.kind = "ACCESSORY"

        parsed = parse_price(item.price_text)
        if parsed is not None:
            item.price_min, item.price_max, item.suspicious_range = parsed
            item.chosen_price = item.price_min if item.suspicious_range else item.price_max

    return items, attached_continuations, skipped


def build_target_rows(items: list[SourceItem]) -> tuple[list[dict[str, object]], list[tuple[int, str, str]], int]:
    skipped: list[tuple[int, str, str]] = []
    glass_rows: list[SourceItem] = []
    accessories: list[SourceItem] = []

    for item in items:
        if item.price_min is None:
            skipped.append((item.excel_row, item.name, "missing price"))
        elif item.kind == "GLASS":
            missing = []
            if item.thickness is None:
                missing.append("thickness")
            if item.width is None or item.height is None:
                missing.append("sheet size")
            if missing:
                skipped.append((item.excel_row, item.name, "missing " + ", ".join(missing)))
            else:
                glass_rows.append(item)
        else:
            accessories.append(item)

    # One glass product per type + colour/finish + thickness. The largest sheet
    # is the one whose sheet price becomes the RETAIL price per m2.
    glass_groups: dict[tuple[str, str, float], list[SourceItem]] = defaultdict(list)
    for item in glass_rows:
        glass_groups[(item.glass_type, item.variant.casefold(), item.thickness)].append(item)

    chosen_glass: dict[int, SourceItem] = {}
    for group in glass_groups.values():
        chosen = max(group, key=lambda candidate: (candidate.width * candidate.height, candidate.excel_row))
        chosen.continuations = [text for member in group for text in member.continuations]
        chosen_glass[id(chosen)] = chosen

    accessory_groups: dict[str, list[SourceItem]] = defaultdict(list)
    for item in accessories:
        accessory_groups[item.name.casefold()].append(item)
    accessory_index = {
        id(item): index
        for group in accessory_groups.values()
        for index, item in enumerate(group)
    }

    target_rows: list[dict[str, object]] = []
    for item in items:
        if item.price_min is None:
            continue
        if item.kind == "GLASS":
            chosen = chosen_glass.get(id(item))
            if chosen is None:
                continue
            group_size = len(glass_groups[(chosen.glass_type, chosen.variant.casefold(), chosen.thickness)])
            target_rows.append({
                "Type": "GLASS",
                "Code": None,
                "Glass type": chosen.glass_type,
                "Colour / finish / pattern": chosen.variant or None,
                "Thickness (mm)": chosen.thickness,
                "Accessory name": None,
                "Accessory unit": None,
                "Sheet width (mm)": chosen.width,
                "Sheet height (mm)": chosen.height,
                "Sheet price (RWF)": chosen.chosen_price,
                "Unit price (RWF)": None,
                "Reorder level (m2)": None,
                "Max level (m2)": None,
                "Notes": note_text(chosen, group_size),
            })
        else:
            group = accessory_groups[item.name.casefold()]
            target_name = accessory_name(item.name, item, accessory_index[id(item)], len(group))
            item.target_name = target_name
            item.unit = accessory_unit(item.name)
            target_rows.append({
                "Type": "ACCESSORY",
                "Code": None,
                "Glass type": None,
                "Colour / finish / pattern": None,
                "Thickness (mm)": None,
                "Accessory name": target_name,
                "Accessory unit": item.unit,
                "Sheet width (mm)": None,
                "Sheet height (mm)": None,
                "Sheet price (RWF)": None,
                "Unit price (RWF)": item.chosen_price,
                "Reorder level (m2)": None,
                "Max level (m2)": None,
                "Notes": note_text(item),
            })

    return target_rows, skipped, len(glass_rows)


def validate_rows(rows: list[dict[str, object]]) -> None:
    names: set[str] = set()
    for row in rows:
        notes = str(row.get("Notes") or "")
        if len(notes) > 255:
            raise SystemExit(f"Notes exceed 255 characters: {notes[:60]}...")
        if row["Type"] == "GLASS":
            if row["Glass type"] not in GLASS_TYPES:
                raise SystemExit(f"Invalid glass type: {row['Glass type']}")
            if not (1 <= float(row["Thickness (mm)"]) <= 50):
                raise SystemExit(f"Invalid thickness: {row['Thickness (mm)']}")
            if not (1 <= int(row["Sheet width (mm)"]) <= 10_000 and 1 <= int(row["Sheet height (mm)"]) <= 10_000):
                raise SystemExit(f"Invalid sheet size: {row['Sheet width (mm)']} x {row['Sheet height (mm)']}")
            if int(row["Sheet price (RWF)"]) <= 0:
                raise SystemExit("Sheet price must be positive")
        else:
            name = str(row["Accessory name"] or "")
            if not name or len(name) > 60:
                raise SystemExit(f"Invalid accessory name: {name!r}")
            if name.casefold() in names:
                raise SystemExit(f"Duplicate accessory name: {name}")
            names.add(name.casefold())
            if row["Accessory unit"] not in ACCESSORY_UNITS:
                raise SystemExit(f"Invalid accessory unit: {row['Accessory unit']}")
            if int(row["Unit price (RWF)"]) <= 0:
                raise SystemExit("Unit price must be positive")


def find_header(ws) -> tuple[int, dict[str, int]]:
    for row in range(1, min(ws.max_row, 25) + 1):
        values = {
            clean_text(ws.cell(row=row, column=column).value): column
            for column in range(1, ws.max_column + 1)
            if clean_text(ws.cell(row=row, column=column).value)
        }
        if all(column in values for column in TARGET_COLUMNS):
            return row, {column: values[column] for column in TARGET_COLUMNS}
    raise SystemExit("The target workbook does not contain the iWarehouse import headers.")


def fill_template(template: Path, output: Path, rows: list[dict[str, object]]) -> int:
    if output.resolve() == template.resolve():
        raise SystemExit("Refusing to overwrite the template. Choose a different output path.")

    output.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template, output)
    workbook = openpyxl.load_workbook(output)
    worksheet = workbook[workbook.sheetnames[0]]
    header_row, columns = find_header(worksheet)

    first_row = worksheet.max_row + 1
    for row in range(header_row + 1, worksheet.max_row + 1):
        if all(not clean_text(worksheet.cell(row=row, column=columns[name]).value) for name in TARGET_COLUMNS):
            first_row = row
            break

    for offset, target_row in enumerate(rows):
        for name in TARGET_COLUMNS:
            value = target_row.get(name)
            worksheet.cell(row=first_row + offset, column=columns[name], value=value)

    workbook.save(output)
    return first_row


def main() -> None:
    parser = argparse.ArgumentParser(description="Convert High Bytes products/prices into the iWarehouse import workbook.")
    parser.add_argument("--source", type=Path, default=SOURCE_WORKBOOK)
    parser.add_argument("--template", type=Path, default=TARGET_TEMPLATE)
    parser.add_argument("--output", type=Path, default=OUTPUT_WORKBOOK)
    args = parser.parse_args()

    if not args.source.exists():
        raise SystemExit(f"Source workbook not found: {args.source}")
    if not args.template.exists():
        raise SystemExit(f"Target template not found: {args.template}")

    items, attached_continuations, read_skips = read_source(args.source)
    target_rows, build_skips, glass_sheet_rows = build_target_rows(items)
    skipped = read_skips + build_skips
    validate_rows(target_rows)

    first_data_row = fill_template(args.template, args.output, target_rows)

    glass_products = sum(1 for row in target_rows if row["Type"] == "GLASS")
    accessories = sum(1 for row in target_rows if row["Type"] == "ACCESSORY")
    suspicious = sum(1 for item in items if item.suspicious_range)

    print("iWarehouse import workbook prepared")
    print(f"  Source rows read:        {len(items)}")
    print(f"  Glass sheet rows:        {glass_sheet_rows}")
    print(f"  Glass products:          {glass_products}")
    print(f"  Duplicate sheets merged: {glass_sheet_rows - glass_products}")
    print(f"  Accessories:             {accessories}")
    print(f"  Products injected:       {len(target_rows)}")
    print(f"  First data row:          {first_data_row}")
    print(f"  Continuation rows:       {attached_continuations}")
    print(f"  Suspicious price ranges: {suspicious}")
    print(f"  Rows skipped:            {len(skipped)}")
    if skipped:
        print("  Skipped source rows:")
        for excel_row, name, reason in skipped:
            print(f"    row {excel_row}: {name} - {reason}")
    print(f"  Output:                  {args.output}")


if __name__ == "__main__":
    main()
